# CUI // SP-CTI
"""dwr-fid-02 — word geometry: the boxes, the bounds, and the seven empties.

The behaviours pinned here are the ones that fail GREEN if they regress: a
status vocabulary that collapses, a rate that reports 0 for a measurement
nobody took, an offset that defaults to 0, a second copy of the DDL, and an
import of the library requirements.txt refuses to declare.
"""
from __future__ import annotations

import ast
import pathlib
import sys

import pytest

sys.path.insert(0, str(pathlib.Path(__file__).parent))

from pdf_fixture import make_pdf  # noqa: E402

from tools.document_intelligence import page_geometry as pg  # noqa: E402

MODULE_PATH = pathlib.Path(pg.__file__)


@pytest.fixture
def two_page_pdf(tmp_path: pathlib.Path) -> pathlib.Path:
    p = tmp_path / "sample.pdf"
    p.write_bytes(
        make_pdf(
            [
                [("Hello world", 72, 700), ("second line here", 72, 680)],
                [("page two words", 100, 600)],
            ]
        )
    )
    return p


@pytest.fixture
def sqlite_conn(tmp_path: pathlib.Path):
    """A throwaway database. NEVER a bare get_connection() — that writes the
    main checkout's data/icdev.db."""
    from tools.db.storage import get_connection

    conn = get_connection(db_path=str(tmp_path / "geom.db"))
    pg.ensure_schema(conn)
    yield conn
    conn.close()


# --------------------------------------------------------------------------- #
# The library prohibition — the card's own rule, read from the source
# --------------------------------------------------------------------------- #
class TestNoPyMuPDF:
    """requirements.txt:218-220 refuses to declare pymupdf (AGPL/commercial).

    pymupdf IS installed on the development host, so an accidental import would
    pass every behavioural test here and produce nothing on a clean install.
    Only the AST can see the difference, so only the AST is asked.
    """

    def test_module_never_imports_pymupdf(self):
        tree = ast.parse(MODULE_PATH.read_text(encoding="utf-8"))
        banned = {"fitz", "pymupdf"}
        found = []
        for node in ast.walk(tree):
            if isinstance(node, ast.Import):
                found += [a.name for a in node.names if a.name.split(".")[0] in banned]
            elif isinstance(node, ast.ImportFrom) and node.module:
                if node.module.split(".")[0] in banned:
                    found.append(node.module)
        assert found == [], f"page_geometry must never import pymupdf; found {found}"

    def test_pdf_path_uses_pdfplumber(self):
        src = MODULE_PATH.read_text(encoding="utf-8")
        assert "import pdfplumber" in src


# --------------------------------------------------------------------------- #
# Extraction and boxes
# --------------------------------------------------------------------------- #
class TestPdfWordBoxes:
    def test_words_carry_a_complete_box_and_page(self, two_page_pdf):
        result = pg.extract_pdf_words(two_page_pdf)
        assert result.status == "extracted"
        assert result.kind == pg.KIND_PDF_WORD_BOX
        assert [w.text for w in result.words if w.page == 1] == [
            "Hello", "world", "second", "line", "here",
        ]
        for w in result.words:
            assert w.x1 > w.x0, "a word with no width is not a placeable box"
            assert w.bottom > w.top, "top is measured from the PAGE TOP, so bottom > top"

    def test_page_dimensions_are_recorded_per_page(self, two_page_pdf):
        result = pg.extract_pdf_words(two_page_pdf)
        assert [p.page for p in result.pages] == [1, 2]
        assert all(p.width == 612.0 and p.height == 792.0 for p in result.pages)
        # The per-page word count is the renderer's sanity check against the
        # rows it fetches, so it must be the real one.
        assert {p.page: p.words for p in result.pages} == {1: 5, 2: 3}

    def test_top_is_measured_from_the_page_top(self, tmp_path):
        """The fixture places text at y=700 from the BOTTOM of a 792pt page.

        pdfplumber's `top` must therefore be near 792-700=92, not near 700. A
        renderer positions with CSS `top`, so a flipped axis draws every word
        upside-down on the page while every count still looks right.
        """
        p = tmp_path / "one.pdf"
        p.write_bytes(make_pdf([[("Anchor", 72, 700)]]))
        word = pg.extract_pdf_words(p).words[0]
        assert 80 < word.top < 105, f"top={word.top} is not measured from the page top"


class TestDocxRuns:
    def test_runs_keep_style_and_distinguish_inherit_from_false(self, tmp_path):
        # A hard import, never importorskip: python-docx is DECLARED in
        # requirements.txt (:155) as a dependency with no fallback, so a skip
        # here would let the whole DOCX half of this feature go untested on the
        # one runner where the package had silently gone missing.
        import docx

        path = tmp_path / "t.docx"
        doc = docx.Document()
        doc.add_heading("Wireless SOP", 1)
        para = doc.add_paragraph("plain then ")
        para.add_run("bold bit").bold = True
        doc.save(str(path))

        result = pg.extract_docx_runs(path)
        assert result.status == "extracted"
        assert result.kind == pg.KIND_DOCX_RUN
        assert result.words == [], "a DOCX has no word boxes, ever"
        styles = {r.text: r.style for r in result.runs}
        assert styles["Wireless SOP"] == "Heading 1"
        bolds = {r.text: r.bold for r in result.runs}
        # None is python-docx's "inherit from the style" and is NOT False: a
        # run that inherits bold from a Heading style is not a run declared
        # non-bold, and squashing the two loses the document's formatting.
        assert bolds["bold bit"] is True
        assert bolds["plain then "] is None


# --------------------------------------------------------------------------- #
# The seven empties
# --------------------------------------------------------------------------- #
class TestStatusVocabulary:
    def test_every_declared_status_is_distinct(self):
        assert len(set(pg.STATUSES)) == len(pg.STATUSES)
        assert pg.RENDERABLE_STATUSES == {"extracted", "truncated"}

    def test_unsupported_format_is_not_a_finding_about_the_document(self, tmp_path):
        txt = tmp_path / "notes.txt"
        txt.write_text("plain text", encoding="utf-8")
        result = pg.capture(txt)
        assert result.status == "unsupported_format"
        assert result.kind is None

    def test_missing_source_is_source_unreadable(self, tmp_path):
        result = pg.capture(tmp_path / "gone.pdf")
        assert result.status == "source_unreadable"

    def test_switched_off_says_so_rather_than_reading_empty(self, two_page_pdf, monkeypatch):
        monkeypatch.setenv(pg.ENABLED_ENV, "0")
        result = pg.capture(two_page_pdf)
        assert result.status == "disabled_by_env"
        assert result.words == []
        # The distinction the switch exists for: this is NOT `no_text_layer`.
        assert result.status != "no_text_layer"

    def test_a_pdf_with_no_text_is_a_measured_zero(self, tmp_path):
        blank = tmp_path / "blank.pdf"
        blank.write_bytes(make_pdf([[]]))
        result = pg.capture(blank, content_type="application/pdf")
        assert result.status == "no_text_layer"
        # It still knows how many pages it read — an unmeasured document does not.
        assert result.pages_total == 1


# --------------------------------------------------------------------------- #
# The bound
# --------------------------------------------------------------------------- #
class TestBoundsAreRealAndReported:
    def test_page_budget_truncates_and_names_the_bound(self, tmp_path):
        p = tmp_path / "many.pdf"
        p.write_bytes(make_pdf([[(f"page{i}", 72, 700)] for i in range(6)]))
        result = pg.extract_pdf_words(p, page_budget=2)
        assert result.status == "truncated"
        assert result.pages_extracted == 2
        assert result.pages_total == 6, "pages_total is the FILE's, not the budget's"
        assert "2 of 6" in result.truncated_reason
        assert {w.page for w in result.words} == {1, 2}

    def test_word_budget_truncates(self, tmp_path):
        p = tmp_path / "wordy.pdf"
        p.write_bytes(make_pdf([[("one two three four five six", 72, 700)]]))
        result = pg.extract_pdf_words(p, word_budget=3)
        assert result.status == "truncated"
        assert len(result.words) == 3
        assert "word budget" in result.truncated_reason

    def test_env_overrides_the_bounds(self, monkeypatch):
        monkeypatch.setenv(pg.MAX_PAGES_ENV, "7")
        assert pg.max_pages() == 7
        # A bound that cannot be honoured falls back rather than disabling
        # capture: max_pages=0 would silently extract nothing.
        monkeypatch.setenv(pg.MAX_PAGES_ENV, "0")
        assert pg.max_pages() == pg.DEFAULT_MAX_PAGES
        monkeypatch.setenv(pg.MAX_PAGES_ENV, "not a number")
        assert pg.max_pages() == pg.DEFAULT_MAX_PAGES


# --------------------------------------------------------------------------- #
# Char offsets
# --------------------------------------------------------------------------- #
class TestCharOffsets:
    def test_offsets_index_the_document_text_that_was_supplied(self, two_page_pdf):
        text = "\n--- Page 1 ---\nHello world\nsecond line here\n--- Page 2 ---\npage two words"
        result = pg.capture(two_page_pdf, content_type="application/pdf", document_text=text)
        assert result.char_basis == "document_text"
        assert result.words_aligned == result.words_total
        for w in result.words:
            assert text[w.char_start:w.char_end] == w.text

    def test_an_unplaceable_word_is_null_never_zero(self, two_page_pdf):
        # Text that contains only some of the words. The rest must come back
        # NULL: an offset of 0 would point them at the first character.
        text = "\n--- Page 1 ---\nHello world\n--- Page 2 ---\nnothing here"
        result = pg.capture(two_page_pdf, content_type="application/pdf", document_text=text)
        placed = [w for w in result.words if w.char_start is not None]
        unplaced = [w for w in result.words if w.char_start is None]
        assert placed and unplaced, "this fixture must exercise both outcomes"
        assert all(w.char_end is None for w in unplaced)
        assert not any(w.char_start == 0 and w.text != "Hello" for w in placed)

    def test_a_changed_text_withholds_offsets_but_keeps_the_boxes(self, two_page_pdf):
        result = pg.capture(
            two_page_pdf,
            content_type="application/pdf",
            document_text="Hello world",
            expected_text_sha256="not-the-hash-this-document-recorded",
        )
        assert result.char_basis == "text_changed"
        assert all(w.char_start is None for w in result.words)
        # The boxes are unaffected: where a word sits on the page does not
        # depend on which library read the text.
        assert result.words and all(w.x1 > w.x0 for w in result.words)

    def test_hash_matches_the_ingest_writer_byte_for_byte(self):
        from tools.document_intelligence.ingest_orchestrator import _sha256

        for sample in ["plain", "unicode — em dash", "\ud800 lone surrogate"]:
            assert pg._sha256_text(sample) == _sha256(sample), (
                "content_sha256 is written by ingest_orchestrator._sha256; a "
                "hash computed under another rule reports text_changed falsely"
            )

    def test_use_text_flow_is_on(self):
        """Measured 14.5% -> 100.0% on a two-column PDF. See the constant."""
        assert pg.USE_TEXT_FLOW is True


class TestAlignRate:
    """A rate that never fabricates, in either direction."""

    @pytest.mark.parametrize(
        "basis,aligned,total,expected",
        [
            # Alignment never ran: None, NOT 0.0.
            ("text_changed", 0, 3347, None),
            ("not_attempted", 0, 0, None),
            ("not_attempted", 0, 500, None),
            # It ran and placed none: a real, measured 0.
            ("unaligned", 0, 500, 0.0),
            # It ran and placed everything.
            ("document_text", 9178, 9178, 100.0),
            # 3346/3347 rounds to 100.0 at one decimal. It must not.
            ("document_text", 3346, 3347, 99.9),
        ],
    )
    def test_rate(self, basis, aligned, total, expected):
        result = pg.GeometryResult(
            status="extracted", kind=pg.KIND_PDF_WORD_BOX,
            char_basis=basis, words_aligned=aligned, words_total=total,
        )
        assert result.align_rate == expected


# --------------------------------------------------------------------------- #
# Persistence
# --------------------------------------------------------------------------- #
class TestPersistence:
    def test_round_trip_and_geometry_row(self, sqlite_conn, two_page_pdf):
        text = "\n--- Page 1 ---\nHello world\nsecond line here\n--- Page 2 ---\npage two words"
        out = pg.capture_and_persist(
            sqlite_conn, "doc-1", two_page_pdf,
            content_type="application/pdf", document_text=text,
            tenant_id="t1", classification="CUI",
        )
        assert out["persisted"] is True

        row = pg.geometry_row(sqlite_conn, "doc-1")
        assert row["status"] == "extracted"
        assert row["kind"] == pg.KIND_PDF_WORD_BOX
        assert row["unit_count"] == 8
        assert row["renderable"] is True
        assert row["align_rate_pct"] == 100.0
        assert [p["page"] for p in row["pages"]] == [1, 2]

        words = pg.page_words(sqlite_conn, "doc-1", 1)
        assert [w["text"] for w in words] == ["Hello", "world", "second", "line", "here"]
        assert words == sorted(words, key=lambda w: w["word_index"])

    def test_reingest_replaces_rather_than_duplicating(self, sqlite_conn, two_page_pdf):
        for _ in range(3):
            pg.capture_and_persist(
                sqlite_conn, "doc-1", two_page_pdf, content_type="application/pdf"
            )
        assert len(pg.page_words(sqlite_conn, "doc-1", 1)) == 5
        cur = sqlite_conn.cursor()
        cur.execute(f"SELECT COUNT(*) FROM {pg.GEOMETRY_TABLE} WHERE doc_id = %s", ("doc-1",))
        assert list(cur.fetchone())[0] == 1

    def test_a_document_with_no_geometry_still_gets_a_row_saying_why(
        self, sqlite_conn, tmp_path
    ):
        txt = tmp_path / "notes.txt"
        txt.write_text("plain", encoding="utf-8")
        pg.capture_and_persist(sqlite_conn, "doc-txt", txt)
        row = pg.geometry_row(sqlite_conn, "doc-txt")
        # The row EXISTS. Absence would be indistinguishable from "nobody
        # looked", which is the one thing the status column is for.
        assert row is not None
        assert row["status"] == "unsupported_format"
        assert row["renderable"] is False
        assert row["align_rate_pct"] is None

    def test_capture_and_persist_never_raises_on_a_broken_connection(self, two_page_pdf):
        class Exploding:
            def cursor(self):
                raise RuntimeError("database is gone")

            def execute(self, *a, **k):
                raise RuntimeError("database is gone")

        out = pg.capture_and_persist(Exploding(), "doc-x", two_page_pdf,
                                     content_type="application/pdf")
        assert out["persisted"] is False
        assert "persist_error" in out


class TestSurvey:
    def test_empty_board_is_unmeasured_not_clean(self, sqlite_conn):
        report = pg.survey(sqlite_conn)
        assert report["state"] in {"unmeasured", "no_geometry_recorded"}
        assert report["state"] != "measured"

    def test_absent_tables_report_the_migration(self, tmp_path):
        from tools.db.storage import get_connection

        conn = get_connection(db_path=str(tmp_path / "bare.db"))
        try:
            report = pg.survey(conn)
            assert report["state"] == "unmeasured"
            assert report["tables_present"] is False
            assert "migrate" in report["reason"]
        finally:
            conn.close()


# --------------------------------------------------------------------------- #
# Structure — the two stories stay apart, the DDL has one copy
# --------------------------------------------------------------------------- #
class TestSchemaShape:
    def _ddl_for(self, table: str) -> str:
        return next(s for s in pg.DDL if f"CREATE TABLE IF NOT EXISTS {table}" in s)

    def _columns(self, table: str) -> set[str]:
        """Declared COLUMN NAMES, with the comments stripped.

        Matching raw substrings against the DDL reads the prose too — these
        tables carry a paragraph each explaining why a column is absent, and
        the word would satisfy a naive `in` check for the very column the
        comment says is not there.
        """
        names = set()
        for line in self._ddl_for(table).splitlines():
            line = line.strip()
            if not line or line.startswith("--") or line.upper().startswith(
                ("CREATE ", ")")
            ):
                continue
            names.add(line.split()[0].rstrip(","))
        return names

    def test_runs_table_has_no_page_column(self):
        """A DOCX has no pages until something renders it.

        A `page` column here would be filled with NULL and read as a
        measurement we failed to take rather than one that cannot exist.
        """
        cols = self._columns(pg.RUNS_TABLE)
        assert "page" not in cols
        assert {"para_index", "run_index", "style"} <= cols
        # And no box columns either: python-docx cannot know them.
        assert not ({"x0", "x1", "top", "bottom"} & cols)

    def test_words_table_has_no_style_column(self):
        """And the converse: a PDF word has a box, not a paragraph style."""
        cols = self._columns(pg.WORDS_TABLE)
        assert "style" not in cols
        assert {"x0", "x1", "top", "bottom", "page", "word_index",
                "char_start", "char_end"} <= cols

    def test_rls_columns_present_on_every_table(self):
        """dic_* is an RLS-owned prefix; a table without these raises on every
        read once the global predicate is attached."""
        for table in (pg.WORDS_TABLE, pg.RUNS_TABLE, pg.GEOMETRY_TABLE):
            assert {"tenant_id", "classification"} <= self._columns(table)

    def test_migration_executes_this_module_ddl_not_a_copy(self):
        """One copy of a table shape, or the two drift at the first new column."""
        migration = pathlib.Path(
            "tools/db/migrations/20260908091858_dic_page_geometry/up.py"
        )
        src = migration.read_text(encoding="utf-8")
        tree = ast.parse(src)
        imported = [
            alias.name
            for node in ast.walk(tree)
            if isinstance(node, ast.ImportFrom)
            and (node.module or "").endswith("page_geometry")
            for alias in node.names
        ]
        assert "DDL" in imported, "the migration must import the DDL, never respell it"

        # No CREATE TABLE in the migration's CODE. Checked against string
        # LITERALS rather than the file text: the module docstring explains why
        # the DDL is plain `CREATE TABLE IF NOT EXISTS`, and a raw `in src`
        # check fails on that explanation while a real second copy of the
        # schema would live in a string the runner executes.
        # Docstrings are identified STRUCTURALLY — the first statement of a
        # module, function or class — not by comparing text to
        # ast.get_docstring(), which returns a cleaned copy that never equals
        # the raw Constant it came from.
        docstring_nodes = set()
        for node in ast.walk(tree):
            if isinstance(node, (ast.Module, ast.FunctionDef, ast.AsyncFunctionDef,
                                 ast.ClassDef)):
                first = (node.body or [None])[0]
                if (
                    isinstance(first, ast.Expr)
                    and isinstance(first.value, ast.Constant)
                    and isinstance(first.value.value, str)
                ):
                    docstring_nodes.add(id(first.value))

        offenders = [
            node.value
            for node in ast.walk(tree)
            if isinstance(node, ast.Constant)
            and isinstance(node.value, str)
            and id(node) not in docstring_nodes
            and "CREATE TABLE" in node.value.upper()
        ]
        assert offenders == [], f"the migration carries its own DDL: {offenders}"

    def test_ingest_schema_executes_the_same_tuple(self):
        """A database the migration has not reached still gets the tables."""
        src = pathlib.Path(
            "tools/document_intelligence/ingest_orchestrator.py"
        ).read_text(encoding="utf-8")
        assert "from tools.document_intelligence.page_geometry import DDL" in src
